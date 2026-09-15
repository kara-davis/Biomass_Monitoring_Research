"""
Score every segmentation-parameter combo produced by sweep_segmentation.R
against field GT heights (Biomass Calcs sheet), using nearest-neighbor
matching logic. Compatible with the updated sweep_segmentation.R output
filename format:

    lmf_fixed_0p25m_k3_ws2.0.csv
    lmf_var_0p25m_k3_var_mid.csv
    li2012_li_default.csv

Writes an Excel workbook with four sheets:
  1. All Results        — every combo sorted by combined score
  2. By Algorithm       — best combo per algorithm family
  3. By CHM Resolution  — best combo per CHM resolution
  4. LMF vs li2012      — head-to-head comparison of best lmf vs li2012

Read-only with respect to existing project files.
"""

import glob
import os
import re
import sys

import numpy as np
import pandas as pd
from pyproj import Transformer
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
SCRIPT_DIR        = os.path.dirname(os.path.abspath(__file__))
REPO_DIR          = os.path.dirname(SCRIPT_DIR)
ARBORETUM_XLSX = r"C:\Users\kdavis99\OneDrive - Cal Poly\Tree Biomass Estimation Research - Documents\GitHub_Repository\Biomass_Monitoring_Research\data\WorkingTrees_Arboretum_Data.xlsx"
BIOMASS_SHEET     = "Biomass Calcs"
TREEPLOTTER_SHEET = "TreePlotter Data"

SWEEP_SUBDIR = sys.argv[1] if len(sys.argv) >= 2 else "sweep_outputs"
OUTPUT_NAME  = sys.argv[2] if len(sys.argv) >= 3 else "sweep_results_summary.xlsx"
SWEEP_DIR    = os.path.join(SCRIPT_DIR, SWEEP_SUBDIR)
OUTPUT_XLSX  = os.path.join(SCRIPT_DIR, OUTPUT_NAME)

MAX_MATCH_DISTANCE_M = 5.0

# ---------------------------------------------------------------------------
# Scoring functions
# ---------------------------------------------------------------------------
def base_score(n_matched_gt, mae):
    """Rewards more matches and lower MAE equally."""
    if n_matched_gt == 0 or (isinstance(mae, float) and np.isnan(mae)):
        return -np.inf
    return n_matched_gt / mae


def combined_score(b_score, abs_count_error_pct):
    """Penalises combos that are far off on total tree count."""
    if not np.isfinite(b_score):
        return -np.inf
    return b_score / (1.0 + abs_count_error_pct)


# ---------------------------------------------------------------------------
# Parse new-format filenames into metadata fields
#
# Patterns:
#   lmf_fixed_<res>_k<k>_ws<ws>.csv
#   lmf_var_<res>_k<k>_<varname>.csv
#   li2012_<configname>.csv
# ---------------------------------------------------------------------------
def parse_filename(fname):
    """Return a dict of metadata extracted from the CSV filename."""
    base = os.path.splitext(os.path.basename(fname))[0]

    # lmf_fixed_0p25m_k3_ws2.0
    m = re.match(r"lmf_fixed_(\w+)_k(\d+)_ws([\d.]+)", base)
    if m:
        return {
            "algorithm_family": "lmf_fixed",
            "chm_resolution":   m.group(1).replace("p", "."),
            "kernel":           int(m.group(2)),
            "ws_setting":       f"fixed_ws{m.group(3)}",
            "config_name":      base,
        }

    # lmf_var_0p25m_k3_var_mid
    m = re.match(r"lmf_var_(\w+)_k(\d+)_(var_\w+)", base)
    if m:
        return {
            "algorithm_family": "lmf_variable",
            "chm_resolution":   m.group(1).replace("p", "."),
            "kernel":           int(m.group(2)),
            "ws_setting":       m.group(3),
            "config_name":      base,
        }

    # li2012_li_default  (no CHM resolution — point-cloud based)
    m = re.match(r"li2012_(.+)", base)
    if m:
        return {
            "algorithm_family": "li2012",
            "chm_resolution":   "N/A (point cloud)",
            "kernel":           None,
            "ws_setting":       m.group(1),
            "config_name":      base,
        }

    # Fallback — unknown format, keep the filename as the label
    return {
        "algorithm_family": "unknown",
        "chm_resolution":   "unknown",
        "kernel":           None,
        "ws_setting":       base,
        "config_name":      base,
    }


# ---------------------------------------------------------------------------
# Field data loading
# ---------------------------------------------------------------------------
def count_treeplotter_total():
    tp = pd.read_excel(ARBORETUM_XLSX, sheet_name=TREEPLOTTER_SHEET)
    tp = tp.dropna(subset=["Latitude", "Longitude"])
    return len(tp)


def load_field_with_gt():
    field = pd.read_excel(ARBORETUM_XLSX, sheet_name=BIOMASS_SHEET)
    field = field.rename(columns={
        "Spec.":      "Species",
        "Tree":       "Tree #",
        "Stem":       "Stem #",
        "Lat":        "lat",
        "Long":       "lon",
        "WT Height":  "WT_height_m",
        "GT Height":  "GT_height_m",
    })
    field = field.dropna(subset=["Species", "Tree #", "Stem #", "Footprint", "lat", "lon"])
    field = field.drop_duplicates(
        subset=["Species", "Tree #", "Stem #", "Footprint"]
    ).reset_index(drop=True)

    transformer = Transformer.from_crs("EPSG:4326", "EPSG:32610", always_xy=True)
    fx, fy = transformer.transform(field["lon"].values, field["lat"].values)
    field["field_x"] = fx
    field["field_y"] = fy
    return field


# ---------------------------------------------------------------------------
# Nearest-neighbour matching (unchanged from original)
# ---------------------------------------------------------------------------
def match_and_score(seg, field):
    all_matches = []
    n_seg = len(seg)

    for fp_id, group in field.groupby("Footprint"):
        group   = group.reset_index(drop=True)
        n_field = len(group)

        sx = seg["seg_x"].values[:, None]
        sy = seg["seg_y"].values[:, None]
        gx = group["field_x"].values[None, :]
        gy = group["field_y"].values[None, :]
        dist = np.sqrt((sx - gx) ** 2 + (sy - gy) ** 2)

        flat_order = np.dstack(
            np.unravel_index(np.argsort(dist, axis=None), dist.shape)
        )[0]
        seg_used   = np.zeros(n_seg,   dtype=bool)
        field_used = np.zeros(n_field, dtype=bool)

        for i, j in flat_order:
            d = dist[i, j]
            if d > MAX_MATCH_DISTANCE_M:
                break
            if seg_used[i] or field_used[j]:
                continue
            seg_used[i]   = True
            field_used[j] = True
            gt = group.loc[j, "GT_height_m"]
            all_matches.append({
                "distance_m":          d,
                "seg_height_m":        seg.loc[i, "seg_height_m"],
                "GT_height_m":         gt,
                "seg_minus_GT_height_m": (
                    seg.loc[i, "seg_height_m"] - gt if pd.notna(gt) else np.nan
                ),
            })

    matched_df    = pd.DataFrame(all_matches)
    n_matched     = len(matched_df)
    have_gt       = (
        matched_df["seg_minus_GT_height_m"].dropna()
        if n_matched else pd.Series(dtype=float)
    )
    n_matched_gt  = len(have_gt)
    mae           = have_gt.abs().mean()            if n_matched_gt else np.nan
    bias          = have_gt.mean()                  if n_matched_gt else np.nan
    rmse          = np.sqrt((have_gt ** 2).mean())  if n_matched_gt else np.nan
    return n_matched, n_matched_gt, mae, bias, rmse


# ---------------------------------------------------------------------------
# Excel formatting helpers
# ---------------------------------------------------------------------------
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT   = Font(color="FFFFFF", bold=True)
ALT_FILL      = PatternFill("solid", fgColor="D6E4F0")
HIGHLIGHT_FILL = PatternFill("solid", fgColor="E2EFDA")  # light green for best row

def format_sheet(ws):
    """Apply header styling and auto-width to a worksheet."""
    for cell in ws[1]:
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for col_idx, col in enumerate(ws.iter_cols(), start=1):
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    # Alternate row shading
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if row_idx % 2 == 0:
            for cell in row:
                cell.fill = ALT_FILL

def highlight_best_row(ws, score_col_idx):
    """Highlight the row with the highest value in score_col_idx (1-based)."""
    best_val = -np.inf
    best_row = None
    for row in ws.iter_rows(min_row=2):
        val = row[score_col_idx - 1].value
        if val is not None and isinstance(val, (int, float)) and val > best_val:
            best_val = val
            best_row = row
    if best_row:
        for cell in best_row:
            cell.fill = HIGHLIGHT_FILL
            cell.font = Font(bold=True)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    print("Loading field data...")
    field             = load_field_with_gt()
    treeplotter_total = count_treeplotter_total()
    print(f"  {len(field)} field records loaded.")
    print(f"  TreePlotter reference count: {treeplotter_total} trees.\n")

    # Find all CSVs in sweep directory
    all_csvs = sorted(glob.glob(os.path.join(SWEEP_DIR, "*.csv")))
    if not all_csvs:
        print(f"No CSVs found in {SWEEP_DIR}. Check the path and re-run.")
        sys.exit(1)

    print(f"Found {len(all_csvs)} CSV files to score.\n")

    rows = []
    for path in all_csvs:
        meta = parse_filename(path)
        print(f"  Scoring: {meta['config_name']} ...")

        seg = pd.read_csv(path)
        seg = seg.rename(columns={
            "treeID":   "seg_treeID",
            "x":        "seg_x",
            "y":        "seg_y",
            "height_m": "seg_height_m",
        })
        n_seg_total = len(seg)

        n_matched, n_matched_gt, mae, bias, rmse = match_and_score(seg, field)

        count_error         = n_seg_total - treeplotter_total
        count_error_pct     = count_error / treeplotter_total
        abs_count_error_pct = abs(count_error_pct)
        b_score             = base_score(n_matched_gt, mae)
        c_score             = combined_score(b_score, abs_count_error_pct)

        rows.append({
            "algorithm_family":               meta["algorithm_family"],
            "chm_resolution":                 meta["chm_resolution"],
            "kernel":                         meta["kernel"],
            "ws_setting":                     meta["ws_setting"],
            "config_name":                    meta["config_name"],
            "n_segmented_total":              n_seg_total,
            "n_treeplotter_reference":        treeplotter_total,
            "count_error (seg-treeplotter)":  count_error,
            "count_error_pct":                round(count_error_pct, 3),
            "segmentation_bias":              (
                "over" if count_error > 0 else ("under" if count_error < 0 else "exact")
            ),
            "n_matched (<=5m)":               n_matched,
            "n_matched_with_GT":              n_matched_gt,
            "bias_m (seg-GT)":                round(bias, 3) if pd.notna(bias) else np.nan,
            "MAE_m":                          round(mae,  3) if pd.notna(mae)  else np.nan,
            "RMSE_m":                         round(rmse, 3) if pd.notna(rmse) else np.nan,
            "score (matches/MAE)":            round(b_score, 3) if np.isfinite(b_score) else np.nan,
            "combined_score":                 round(c_score, 3) if np.isfinite(c_score) else np.nan,
        })

    all_results = (
        pd.DataFrame(rows)
        .sort_values("combined_score", ascending=False)
        .reset_index(drop=True)
    )

    # ------------------------------------------------------------------
    # Sheet 2: best combo per algorithm family
    # ------------------------------------------------------------------
    by_algorithm = (
        all_results
        .sort_values("combined_score", ascending=False)
        .groupby("algorithm_family", as_index=False)
        .first()
        .sort_values("combined_score", ascending=False)
        .reset_index(drop=True)
    )

    # ------------------------------------------------------------------
    # Sheet 3: best combo per CHM resolution
    # ------------------------------------------------------------------
    by_resolution = (
        all_results[all_results["chm_resolution"] != "N/A (point cloud)"]
        .sort_values("combined_score", ascending=False)
        .groupby("chm_resolution", as_index=False)
        .first()
        .sort_values("combined_score", ascending=False)
        .reset_index(drop=True)
    )

    # ------------------------------------------------------------------
    # Sheet 4: LMF vs li2012 head-to-head
    # Best lmf_fixed, best lmf_variable, and best li2012 side by side
    # ------------------------------------------------------------------
    comparison_rows = []
    for family in ["lmf_fixed", "lmf_variable", "li2012"]:
        subset = all_results[all_results["algorithm_family"] == family]
        if len(subset) == 0:
            continue
        best = subset.iloc[0].copy()
        best["rank_within_family"] = 1
        comparison_rows.append(best)

    lmf_vs_li2012 = pd.DataFrame(comparison_rows).reset_index(drop=True)

    # Add a winner column
    if len(lmf_vs_li2012) > 0:
        best_idx = lmf_vs_li2012["combined_score"].idxmax()
        lmf_vs_li2012["WINNER"] = ""
        lmf_vs_li2012.loc[best_idx, "WINNER"] = "★ BEST"

    # ------------------------------------------------------------------
    # Write Excel workbook
    # ------------------------------------------------------------------
    print(f"\nWriting results to {OUTPUT_XLSX} ...")
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        all_results.to_excel(writer,   sheet_name="All Results",       index=False)
        by_algorithm.to_excel(writer,  sheet_name="By Algorithm",      index=False)
        by_resolution.to_excel(writer, sheet_name="By CHM Resolution", index=False)
        lmf_vs_li2012.to_excel(writer, sheet_name="LMF vs li2012",     index=False)

        wb = writer.book
        score_col = all_results.columns.get_loc("combined_score") + 1

        for sheet_name in ["All Results", "By Algorithm", "By CHM Resolution", "LMF vs li2012"]:
            ws = wb[sheet_name]
            format_sheet(ws)
            highlight_best_row(ws, score_col)

    print(f"Done. Wrote {OUTPUT_XLSX}\n")

    # ------------------------------------------------------------------
    # Console summary
    # ------------------------------------------------------------------
    print("=" * 70)
    print("TOP 10 COMBOS (all algorithms)")
    print("=" * 70)
    top10_cols = [
        "algorithm_family", "chm_resolution", "config_name",
        "n_segmented_total", "n_matched_with_GT", "MAE_m", "combined_score"
    ]
    print(all_results[top10_cols].head(10).to_string(index=False))

    print("\n" + "=" * 70)
    print("BEST PER ALGORITHM FAMILY")
    print("=" * 70)
    print(by_algorithm[top10_cols].to_string(index=False))

    print("\n" + "=" * 70)
    print("LMF vs LI2012 HEAD-TO-HEAD")
    print("=" * 70)
    h2h_cols = [
        "algorithm_family", "config_name",
        "n_segmented_total", "n_matched_with_GT",
        "MAE_m", "combined_score", "WINNER"
    ]
    available = [c for c in h2h_cols if c in lmf_vs_li2012.columns]
    print(lmf_vs_li2012[available].to_string(index=False))


if __name__ == "__main__":
    main()
